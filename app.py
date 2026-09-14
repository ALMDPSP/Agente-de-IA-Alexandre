import base64
import hashlib
import hmac
import io
import logging
import os
import re
import time
import uuid
from datetime import timedelta
from pathlib import Path

import pyotp
import qrcode
from docx import Document
from dotenv import load_dotenv
from flask import Flask, jsonify, redirect, render_template, request, session, url_for
from flask_login import LoginManager, UserMixin, current_user, login_required, login_user, logout_user
from flask_wtf.csrf import CSRFProtect
from openai import OpenAI
from openpyxl import load_workbook
from pypdf import PdfReader

from database import (
    add_chat_message, add_history, clear_chat, create_project, database_configured,
    delete_document, delete_history, delete_project, init_database, insert_document,
    load_state, replace_state, set_active_project, update_project,
)

load_dotenv()

app = Flask(__name__)
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("alexandre-ai")

app.config.update(
    SECRET_KEY=os.getenv("SECRET_KEY", "dev-change-me"),
    REMEMBER_COOKIE_HTTPONLY=True,
    REMEMBER_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
    MAX_CONTENT_LENGTH=15 * 1024 * 1024,
)

if os.getenv("RENDER"):
    app.config["SESSION_COOKIE_SECURE"] = True
    app.config["REMEMBER_COOKIE_SECURE"] = True

csrf = CSRFProtect(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Faça login para continuar."
login_manager.login_message_category = "warning"

try:
    init_database()
except Exception as exc:
    logger.exception("Não foi possível inicializar o PostgreSQL: %s", exc)


SYSTEM_PROMPT = """Você é Alexandre AI, o assistente pessoal de Alexandre.
Responda sempre em português do Brasil, salvo se o usuário pedir outro idioma.
Seja claro, objetivo, técnico quando necessário e útil.

Você pode receber CONTEXTO PRIVADO DE PROJETO com trechos de documentos enviados pelo usuário.
Use esse conteúdo como fonte primária quando a pergunta estiver relacionada ao projeto.
Nunca invente informações ausentes no contexto.
Quando utilizar um documento, mencione naturalmente o nome do arquivo quando isso ajudar.
Se o contexto for insuficiente, diga o que falta.
"""

PREAUTH_TTL_SECONDS = 300
MAX_MFA_ATTEMPTS = 5

SUPPORTED_TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".log", ".xml", ".html", ".htm",
}
SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".xlsm",
}
SUPPORTED_EXTENSIONS = SUPPORTED_TEXT_EXTENSIONS | SUPPORTED_DOCUMENT_EXTENSIONS


class EnvUser(UserMixin):
    def __init__(self):
        self.id = "admin"
        self.name = os.getenv("ADMIN_NAME", "Alexandre")
        self.email = os.getenv("ADMIN_EMAIL", "").strip().lower()
        self.role = "admin"


def get_admin_user():
    return EnvUser()


@login_manager.user_loader
def load_user(user_id):
    if user_id == "admin":
        return get_admin_user()
    return None


def env_bool(name, default=False):
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on", "sim"}


def mfa_enabled():
    return env_bool("MFA_ENABLED", False)


def mfa_setup_enabled():
    return env_bool("MFA_SETUP_ENABLED", False)


def mfa_secret():
    source = app.config["SECRET_KEY"]
    digest = hmac.new(
        str(source).encode("utf-8"),
        b"alexandre-ai-mfa-v1",
        hashlib.sha256,
    ).digest()
    return base64.b32encode(digest).decode("ascii").rstrip("=")


def mfa_totp():
    return pyotp.TOTP(mfa_secret(), digits=6, interval=30)


def mfa_uri():
    user = get_admin_user()
    account_name = user.email or "Alexandre"
    return mfa_totp().provisioning_uri(
        name=account_name,
        issuer_name="Alexandre AI",
    )


def qr_data_uri(text):
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=2,
    )
    qr.add_data(text)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white")

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def start_preauth(remember):
    session["preauth"] = True
    session["preauth_at"] = int(time.time())
    session["remember_after_mfa"] = bool(remember)
    session["mfa_attempts"] = 0


def clear_preauth():
    for key in ("preauth", "preauth_at", "remember_after_mfa", "mfa_attempts"):
        session.pop(key, None)


def preauth_valid():
    if not session.get("preauth"):
        return False

    try:
        age = time.time() - float(session.get("preauth_at", 0))
    except (TypeError, ValueError):
        return False

    return 0 <= age <= PREAUTH_TTL_SECONDS


def finish_login():
    remember = bool(session.get("remember_after_mfa"))
    login_user(get_admin_user(), remember=remember)
    clear_preauth()


def normalize_text(text, max_chars=200000):
    text = (text or "").replace("\x00", "")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()[:max_chars]


def split_text_chunks(text, filename, chunk_size=2400, overlap=350):
    text = normalize_text(text)
    if not text:
        return []

    chunks = []
    cursor = 0
    index = 1

    while cursor < len(text):
        end = min(len(text), cursor + chunk_size)
        piece = text[cursor:end]

        if end < len(text):
            natural_break = max(
                piece.rfind("\n\n"),
                piece.rfind(". "),
                piece.rfind("\n"),
            )
            if natural_break > chunk_size * 0.55:
                end = cursor + natural_break + 1
                piece = text[cursor:end]

        piece = piece.strip()
        if piece:
            chunks.append({
                "id": f"chunk_{index}",
                "index": index,
                "filename": filename,
                "text": piece,
            })

        if end >= len(text):
            break

        cursor = max(cursor + 1, end - overlap)
        index += 1

    return chunks[:180]


def extract_pdf(file_obj):
    reader = PdfReader(file_obj)
    parts = []
    for idx, page in enumerate(reader.pages[:100], start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[Página {idx}]\n{text}")
    return normalize_text("\n\n".join(parts))


def extract_docx(file_obj):
    document = Document(file_obj)
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))

    return normalize_text("\n".join(parts))


def extract_xlsx(file_obj):
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    parts = []

    for sheet in workbook.worksheets[:25]:
        parts.append(f"[Planilha: {sheet.title}]")
        included = 0

        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                parts.append(" | ".join(values))
                included += 1

            if included >= 2500:
                parts.append("[Conteúdo truncado nesta planilha]")
                break

    return normalize_text("\n".join(parts))


def extract_text_file(file_obj):
    data = file_obj.read()
    if isinstance(data, bytes):
        return normalize_text(data.decode("utf-8", errors="replace"))
    return normalize_text(str(data))


def extract_file_content(file_obj, filename):
    suffix = Path(filename or "").suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Formato não suportado. Use PDF, DOCX, XLSX, TXT, MD, CSV, JSON, XML, HTML ou LOG."
        )

    if suffix == ".pdf":
        return extract_pdf(file_obj)
    if suffix == ".docx":
        return extract_docx(file_obj)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(file_obj)

    return extract_text_file(file_obj)


def sanitize_history(raw_history):
    clean = []
    if not isinstance(raw_history, list):
        return clean

    for item in raw_history[-16:]:
        if not isinstance(item, dict):
            continue

        role = item.get("role") or item.get("type")
        content = item.get("content") or item.get("text")

        if role == "assistant":
            normalized_role = "assistant"
        elif role == "user":
            normalized_role = "user"
        else:
            continue

        if not isinstance(content, str):
            continue

        content = content.strip()
        if content:
            clean.append({
                "role": normalized_role,
                "content": content[:12000],
            })

    return clean


def sanitize_project_context(raw_context):
    if not isinstance(raw_context, str):
        return ""
    return normalize_text(raw_context, max_chars=52000)


def provider_configurations():
    providers = []

    groq_key = os.getenv("GROQ_API_KEY")
    if groq_key:
        providers.append({
            "name": "Groq",
            "client": OpenAI(
                api_key=groq_key,
                base_url="https://api.groq.com/openai/v1",
                timeout=35.0,
                max_retries=0,
            ),
            "model": os.getenv("GROQ_MODEL", "openai/gpt-oss-20b"),
        })

    gemini_key = os.getenv("GEMINI_API_KEY")
    if gemini_key:
        providers.append({
            "name": "Gemini",
            "client": OpenAI(
                api_key=gemini_key,
                base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
                timeout=35.0,
                max_retries=0,
            ),
            "model": os.getenv("GEMINI_MODEL", "gemini-3.8-flash"),
        })

    cloudflare_key = os.getenv("CLOUDFLARE_API_TOKEN")
    cloudflare_account = os.getenv("CLOUDFLARE_ACCOUNT_ID")
    if cloudflare_key and cloudflare_account:
        providers.append({
            "name": "Cloudflare",
            "client": OpenAI(
                api_key=cloudflare_key,
                base_url=f"https://api.cloudflare.com/client/v4/accounts/{cloudflare_account}/ai/v1",
                timeout=35.0,
                max_retries=0,
            ),
            "model": os.getenv(
                "CLOUDFLARE_MODEL",
                "@cf/google/gemma-4-26b-a4b-it",
            ),
        })

    return providers


def ask_with_fallback(messages):
    providers = provider_configurations()

    if not providers:
        raise RuntimeError("Nenhum provedor de IA está configurado.")

    failures = []

    for provider in providers:
        try:
            logger.info("Tentando provedor=%s modelo=%s", provider["name"], provider["model"])

            response = provider["client"].chat.completions.create(
                model=provider["model"],
                messages=messages,
                temperature=0.35,
                max_tokens=1800,
            )

            content = response.choices[0].message.content
            if not content:
                raise RuntimeError("Resposta vazia.")

            return {
                "reply": content.strip(),
                "provider": provider["name"],
                "model": provider["model"],
            }

        except Exception as exc:
            logger.warning("Falha no provedor %s: %s", provider["name"], str(exc)[:500])
            failures.append(provider["name"])

    raise RuntimeError(
        "Os provedores configurados não responderam. "
        f"Tentativas: {', '.join(failures)}."
    )


@app.after_request
def set_security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "SAMEORIGIN"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    error = None

    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        remember = request.form.get("remember") == "on"

        admin_email = os.getenv("ADMIN_EMAIL", "").strip().lower()
        admin_password = os.getenv("ADMIN_PASSWORD", "")

        email_ok = bool(admin_email) and hmac.compare_digest(email, admin_email)
        password_ok = bool(admin_password) and hmac.compare_digest(password, admin_password)

        if email_ok and password_ok:
            if mfa_enabled():
                start_preauth(remember)
                return redirect(url_for("mfa_challenge"))

            login_user(get_admin_user(), remember=remember)
            return redirect(url_for("dashboard"))

        error = "E-mail ou senha inválidos."

    return render_template("login.html", error=error, mfa_enabled=mfa_enabled())


@app.route("/mfa", methods=["GET", "POST"])
def mfa_challenge():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if not mfa_enabled():
        return redirect(url_for("login"))

    if not preauth_valid():
        clear_preauth()
        return redirect(url_for("login"))

    error = None

    if request.method == "POST":
        attempts = int(session.get("mfa_attempts", 0))
        if attempts >= MAX_MFA_ATTEMPTS:
            clear_preauth()
            return redirect(url_for("login"))

        code = "".join(ch for ch in request.form.get("code", "") if ch.isdigit())

        if len(code) == 6 and mfa_totp().verify(code, valid_window=1):
            finish_login()
            return redirect(url_for("dashboard"))

        attempts += 1
        session["mfa_attempts"] = attempts

        if attempts >= MAX_MFA_ATTEMPTS:
            clear_preauth()
            return render_template(
                "mfa.html",
                error="Limite de tentativas atingido. Faça o login novamente.",
                setup_enabled=False,
                locked=True,
            )

        error = f"Código inválido. Restam {MAX_MFA_ATTEMPTS - attempts} tentativa(s)."

    return render_template(
        "mfa.html",
        error=error,
        setup_enabled=mfa_setup_enabled(),
        locked=False,
    )


@app.route("/mfa/setup", methods=["GET", "POST"])
def mfa_setup():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if not mfa_enabled() or not mfa_setup_enabled():
        return redirect(url_for("mfa_challenge"))

    if not preauth_valid():
        clear_preauth()
        return redirect(url_for("login"))

    error = None

    if request.method == "POST":
        code = "".join(ch for ch in request.form.get("code", "") if ch.isdigit())

        if len(code) == 6 and mfa_totp().verify(code, valid_window=1):
            finish_login()
            return redirect(url_for("dashboard"))

        error = "Código inválido. Confira o relógio do celular e tente novamente."

    return render_template(
        "mfa_setup.html",
        error=error,
        qr_data=qr_data_uri(mfa_uri()),
        manual_key=mfa_secret(),
        account_email=get_admin_user().email,
    )


@app.route("/dashboard")
@login_required
def dashboard():
    return render_template(
        "dashboard.html",
        user=current_user,
        mfa_enabled=mfa_enabled(),
    )


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    logout_user()
    clear_preauth()
    return redirect(url_for("login"))


@app.route("/api/files/extract", methods=["POST"])
@login_required
def extract_uploaded_file():
    """Mantido para compatibilidade; extrai sem persistir."""
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Selecione um arquivo."}), 400
    try:
        text = extract_file_content(uploaded.stream, uploaded.filename)
        if not text:
            return jsonify({"error": "O arquivo não possui texto legível para importar."}), 400
        chunks = split_text_chunks(text, uploaded.filename)
        return jsonify({
            "title": uploaded.filename,
            "content": text,
            "chunks": chunks,
            "sourceType": "local_file",
            "characters": len(text),
            "chunkCount": len(chunks),
            "size": request.content_length or 0,
        })
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        logger.exception("Erro ao extrair arquivo")
        return jsonify({"error": f"Não foi possível ler o arquivo: {str(exc)[:160]}"}), 400


def _owner_id():
    return str(current_user.get_id())


def _db_error(exc):
    logger.exception("Erro PostgreSQL")
    return jsonify({"error": f"Falha de persistência no PostgreSQL: {str(exc)[:180]}"}), 500


@app.route("/api/state")
@login_required
def api_state():
    if not database_configured():
        return jsonify({"error": "DATABASE_URL não configurada no Render."}), 503
    try:
        return jsonify(load_state(_owner_id()))
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/state/import", methods=["POST"])
@login_required
def api_state_import():
    try:
        replace_state(_owner_id(), request.get_json(silent=True) or {})
        return jsonify({"ok": True})
    except (ValueError, TypeError) as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/projects", methods=["POST"])
@login_required
def api_projects_create():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Informe o nome do projeto."}), 400
    project = {
        "id": str(data.get("id") or f"project_{uuid.uuid4()}"),
        "name": name[:200],
        "description": str(data.get("description") or "")[:10000],
        "status": str(data.get("status") or "Planejamento")[:100],
    }
    try:
        return jsonify(create_project(_owner_id(), project)), 201
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/projects/<project_id>", methods=["PUT", "DELETE"])
@login_required
def api_project_item(project_id):
    try:
        if request.method == "DELETE":
            delete_project(_owner_id(), project_id)
            return jsonify({"ok": True})
        data = request.get_json(silent=True) or {}
        name = str(data.get("name") or "").strip()
        if not name:
            return jsonify({"error": "Informe o nome do projeto."}), 400
        row = update_project(_owner_id(), project_id, {
            "name": name[:200],
            "description": str(data.get("description") or "")[:10000],
            "status": str(data.get("status") or "Planejamento")[:100],
        })
        if not row:
            return jsonify({"error": "Projeto não encontrado."}), 404
        return jsonify({"ok": True, "updatedAt": row["updated_at"].isoformat()})
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/settings/active-project", methods=["PUT"])
@login_required
def api_active_project():
    data = request.get_json(silent=True) or {}
    try:
        set_active_project(_owner_id(), str(data.get("projectId") or ""))
        return jsonify({"ok": True})
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/projects/<project_id>/documents", methods=["POST"])
@login_required
def api_project_document(project_id):
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Selecione um arquivo."}), 400
    try:
        text = extract_file_content(uploaded.stream, uploaded.filename)
        if not text:
            return jsonify({"error": "O arquivo não possui texto legível para importar."}), 400
        chunks = split_text_chunks(text, uploaded.filename)
        document = {
            "id": f"doc_{uuid.uuid4()}",
            "title": uploaded.filename,
            "size": int(request.content_length or 0),
            "characters": len(text),
            "chunks": chunks,
        }
        saved = insert_document(_owner_id(), project_id, document)
        return jsonify(saved), 201
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/documents/<document_id>", methods=["DELETE"])
@login_required
def api_document_delete(document_id):
    try:
        if not delete_document(_owner_id(), document_id):
            return jsonify({"error": "Documento não encontrado."}), 404
        return jsonify({"ok": True})
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/chat", methods=["POST", "DELETE"])
@login_required
def api_chat():
    try:
        if request.method == "DELETE":
            clear_chat(_owner_id())
            return jsonify({"ok": True})
        data = request.get_json(silent=True) or {}
        text = str(data.get("text") or "").strip()
        if not text:
            return jsonify({"error": "Mensagem vazia."}), 400
        row = add_chat_message(_owner_id(), {
            "text": text,
            "type": data.get("type"),
            "provider": data.get("provider") or "",
            "documents": data.get("documents") or [],
        })
        return jsonify({"ok": True, "id": row["id"]}), 201
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/history", methods=["POST", "DELETE"])
@login_required
def api_history():
    try:
        if request.method == "DELETE":
            history_id = request.args.get("id")
            delete_history(_owner_id(), history_id)
            return jsonify({"ok": True})
        data = request.get_json(silent=True) or {}
        query = str(data.get("query") or "").strip()
        if not query:
            return jsonify({"error": "Consulta vazia."}), 400
        item = {
            "id": str(data.get("id") or f"search_{uuid.uuid4()}"),
            "query": query[:20000],
            "projectId": str(data.get("projectId") or ""),
            "projectName": str(data.get("projectName") or "Conversa geral"),
            "provider": str(data.get("provider") or ""),
            "usedDocuments": data.get("usedDocuments") or [],
        }
        add_history(_owner_id(), item)
        return jsonify(item), 201
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/backup/export")
@login_required
def api_backup_export():
    try:
        payload = load_state(_owner_id())
        payload["version"] = 3
        payload["storage"] = "postgresql"
        return jsonify(payload)
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/backup/import", methods=["POST"])
@login_required
def api_backup_import():
    try:
        replace_state(_owner_id(), request.get_json(silent=True) or {})
        return jsonify({"ok": True})
    except (ValueError, TypeError) as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return _db_error(exc)


@app.route("/api/agent/status")
@login_required
def agent_status():
    configured = [
        {"provider": provider["name"], "model": provider["model"]}
        for provider in provider_configurations()
    ]

    return jsonify({
        "configured": configured,
        "fallback_order": [provider["provider"] for provider in configured],
        "mfa": mfa_enabled(),
    })


@app.route("/api/agent", methods=["POST"])
@login_required
def agent():
    data = request.get_json(silent=True) or {}
    message = str(data.get("message", "")).strip()

    if not message:
        return jsonify({"error": "Digite uma mensagem."}), 400

    history = sanitize_history(data.get("history", []))
    project_context = sanitize_project_context(data.get("projectContext", ""))

    messages = [{"role": "system", "content": SYSTEM_PROMPT}]

    if project_context:
        messages.append({
            "role": "system",
            "content": (
                "CONTEXTO PRIVADO DO PROJETO:\n"
                "Os trechos abaixo foram selecionados automaticamente dos arquivos "
                "enviados pelo usuário porque parecem relevantes para a pergunta.\n\n"
                f"{project_context}"
            ),
        })

    messages.extend(history)

    if not history or history[-1].get("role") != "user" or history[-1].get("content") != message:
        messages.append({"role": "user", "content": message})

    try:
        result = ask_with_fallback(messages)
        return jsonify(result)
    except RuntimeError as exc:
        return jsonify({"error": str(exc)}), 503
    except Exception:
        logger.exception("Erro inesperado no Agente IA")
        return jsonify({"error": "Erro interno ao processar a solicitação."}), 500


@app.route("/health")
def health():
    return jsonify({"status": "ok", "database": "configured" if database_configured() else "missing"}), 200


if __name__ == "__main__":
    app.run(debug=True)
